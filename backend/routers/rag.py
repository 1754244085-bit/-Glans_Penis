"""
上传文档、解析文本、建索引（配合 RAG）。
文档问答是独立能力，单独路由便于限流与扩展格式。
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from deps import get_current_user_id  # noqa: E402
import rag_store  # noqa: E402

router = APIRouter()


def _extract_text(file_name: str, content: bytes) -> str:
    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    if ext == "txt":
        return content.decode("utf-8", errors="ignore")
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext in {"docx", "doc"}:
        import docx

        doc = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext in {"xlsx", "xls"}:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"# 工作表: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines)
    raise HTTPException(status_code=400, detail="仅支持 PDF/Word/TXT/Excel")


@router.post("/upload")
async def upload_doc(
    file: UploadFile = File(...),
    user_id: int = Depends(get_current_user_id),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="缺少文件名")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件过大（<=15MB）")
    text = _extract_text(file.filename, data)
    try:
        doc = rag_store.add_document(user_id, file.filename, text)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {
        "ok": True,
        "doc_id": doc.doc_id,
        "file_name": doc.file_name,
        "chunk_count": doc.chunk_count,
    }


@router.get("/docs")
def list_docs(user_id: int = Depends(get_current_user_id)):
    docs = rag_store.list_documents(user_id)
    return {
        "docs": [
            {
                "doc_id": d.doc_id,
                "file_name": d.file_name,
                "chunk_count": d.chunk_count,
                "created_at": d.created_at,
            }
            for d in docs
        ]
    }


@router.delete("/docs/{doc_id}")
def delete_doc(doc_id: str, user_id: int = Depends(get_current_user_id)):
    ok = rag_store.delete_document(user_id, doc_id)
    if not ok:
        raise HTTPException(status_code=404, detail="文档不存在")
    return {"ok": True}

